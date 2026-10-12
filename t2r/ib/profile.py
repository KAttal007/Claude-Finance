"""Company profile slide (design doc 12): summary, key financials, valuation, recent events.

Builds onto a blank 16:9 slide. PHASE 5: switch to the user's branded template
(t2r/ib/templates/profile.pptx) and fill its placeholders instead; add shareholding and a price
chart once prices and shareholding data are loaded.
"""

import io
from dataclasses import dataclass, field

from pptx import Presentation
from pptx.util import Inches, Pt


@dataclass
class ProfileData:
    name: str
    symbol: str
    business_summary: str
    financials: list[tuple[str, list[str]]]  # (row label, values per period)
    periods: list[str]
    valuation: list[tuple[str, str]] = field(default_factory=list)
    events: list[str] = field(default_factory=list)


def build(data: ProfileData) -> bytes:
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank

    def textbox(
        x: float, y: float, w: float, h: float, text: str, size: int = 12, bold: bool = False
    ) -> None:
        tf = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)).text_frame
        tf.word_wrap = True
        tf.text = text
        for p in tf.paragraphs:
            p.font.size, p.font.bold = Pt(size), bold

    textbox(0.4, 0.3, 12.5, 0.6, f"{data.name} ({data.symbol})", 24, True)
    textbox(0.4, 1.0, 6.2, 2.4, data.business_summary, 12)

    rows, cols = len(data.financials) + 1, len(data.periods) + 1
    table = slide.shapes.add_table(
        rows, cols, Inches(6.9), Inches(1.0), Inches(6.0), Inches(0.3 * rows)
    ).table
    table.cell(0, 0).text = "INR crore"
    for c, p in enumerate(data.periods, start=1):
        table.cell(0, c).text = p
    for r, (label, values) in enumerate(data.financials, start=1):
        table.cell(r, 0).text = label
        for c, v in enumerate(values, start=1):
            table.cell(r, c).text = v

    textbox(
        0.4, 3.7, 6.2, 2.0, "Valuation\n" + "\n".join(f"{k}: {v}" for k, v in data.valuation), 12
    )
    textbox(6.9, 4.6, 6.0, 2.4, "Recent events\n" + "\n".join(f"- {e}" for e in data.events), 12)
    out = io.BytesIO()
    prs.save(out)
    return out.getvalue()
