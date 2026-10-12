"""Turn grounded report text into Word and PowerPoint files (design doc 11.1)."""

import io
import re

from docx import Document
from pptx import Presentation
from pptx.util import Inches, Pt

DISCLAIMER = (
    "Analytical summary generated from public filings. Not investment advice; "
    "no rating or target price."
)


def note_docx(title: str, markdown: str) -> bytes:
    doc = Document()
    doc.add_heading(title, level=0)
    for block in markdown.split("\n"):
        line = block.rstrip()
        if not line:
            continue
        if line.startswith("#"):
            level = min(len(line) - len(line.lstrip("#")), 3)
            doc.add_heading(line.lstrip("# "), level=level)
        elif re.match(r"^\s*[-*] ", line):
            doc.add_paragraph(line.lstrip(" -*"), style="List Bullet")
        else:
            doc.add_paragraph(line)
    doc.add_paragraph(DISCLAIMER).runs[0].italic = True
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


def summary_pptx(title: str, slides: list[tuple[str, list[str]]]) -> bytes:
    prs = Presentation()
    first = prs.slides.add_slide(prs.slide_layouts[0])
    first.shapes.title.text = title
    first.placeholders[1].text = DISCLAIMER
    for heading, bullets in slides:
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = heading
        body = slide.placeholders[1].text_frame
        body.clear()
        for i, b in enumerate(bullets):
            para = body.paragraphs[0] if i == 0 else body.add_paragraph()
            para.text = b
            para.font.size = Pt(16)
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    out = io.BytesIO()
    prs.save(out)
    return out.getvalue()
