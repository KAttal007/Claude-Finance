"""Locate periods and rows in an analyst workbook (headless, openpyxl)."""

from datetime import date, datetime

from openpyxl.utils import get_column_letter
from openpyxl.workbook import Workbook
from openpyxl.worksheet.worksheet import Worksheet

from t2r.modelsync.mapping import Mapping


def fiscal_label(period_end: date) -> str:
    """Indian fiscal quarter label: Apr-Jun 2026 -> Q1FY27."""
    fy = period_end.year + 1 if period_end.month >= 4 else period_end.year
    quarter = {6: 1, 9: 2, 12: 3, 3: 4}.get(period_end.month)
    if quarter is None:
        raise ValueError(f"{period_end} is not a fiscal quarter end")
    return f"Q{quarter}FY{fy % 100:02d}"


def find_row(ws: Worksheet, label: str, label_column: int = 1) -> int | None:
    target = label.strip().lower()
    for row in range(1, ws.max_row + 1):
        value = ws.cell(row=row, column=label_column).value
        if isinstance(value, str) and value.strip().lower() == target:
            return row
    return None


class PeriodAxis:
    """Header row of period labels/dates plus an A/E flag row on the mapping's period sheet."""

    def __init__(self, wb: Workbook, mapping: Mapping) -> None:
        axis = mapping.period_axis
        self.ws = wb[axis.sheet]
        self.label_column = axis.label_column
        header = find_row(self.ws, axis.header_row_label, axis.label_column)
        flag = find_row(self.ws, axis.flag_row_label, axis.label_column)
        if header is None or flag is None:
            raise LookupError(f"period axis rows not found on sheet {axis.sheet!r}")
        self.header_row, self.flag_row = header, flag

    def columns(self) -> list[int]:
        cols = []
        for col in range(self.label_column + 1, self.ws.max_column + 1):
            if self.ws.cell(row=self.header_row, column=col).value not in (None, ""):
                cols.append(col)
        return cols

    def label(self, col: int) -> str:
        value = self.ws.cell(row=self.header_row, column=col).value
        if isinstance(value, datetime | date):
            d = value.date() if isinstance(value, datetime) else value
            return fiscal_label(d)
        return str(value).strip().upper().replace(" ", "")

    def flag(self, col: int) -> str:
        return str(self.ws.cell(row=self.flag_row, column=col).value or "").strip().upper()

    def column_for(self, period_end: date) -> int | None:
        wanted = fiscal_label(period_end)
        return next((c for c in self.columns() if self.label(c) == wanted), None)

    def estimate_columns(self) -> list[int]:
        return [c for c in self.columns() if self.flag(c) == "E"]


def cell_ref(col: int, row: int) -> str:
    return f"{get_column_letter(col)}{row}"
