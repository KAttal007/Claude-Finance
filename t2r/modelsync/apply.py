"""Apply approved change items (design doc 10.5).

Headless path: openpyxl writes, optional LibreOffice recalculation. The xlwings path (native Excel
recalculation, formatting preserved) lives in t2r.excel_bridge and takes the same item list.
The original workbook is never overwritten: the caller stores the result as a new model_version.
"""

import io
import shutil
import subprocess
import tempfile
from dataclasses import dataclass
from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl import load_workbook


@dataclass
class ApprovedItem:
    kind: str
    sheet: str
    cell: str
    value: str | None  # final value (edited or as proposed)
    formula: str | None = None


def coerce(value: str | None) -> object:
    if value is None:
        return None
    try:
        d = Decimal(value)
    except InvalidOperation:
        return value
    return int(d) if d == d.to_integral_value() and abs(d) < 10**15 else float(d)


def apply_headless(workbook: bytes, items: list[ApprovedItem]) -> bytes:
    wb = load_workbook(io.BytesIO(workbook))
    for item in items:
        ws = wb[item.sheet]
        ws[item.cell].value = item.formula if item.formula else coerce(item.value)
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def recalc_with_libreoffice(workbook: bytes) -> bytes:
    """Round-trip through headless LibreOffice so cached values are recalculated.

    Returns the input unchanged when soffice is not installed (cached values stay stale until
    the analyst opens the file in Excel).
    """
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if soffice is None:
        return workbook
    with tempfile.TemporaryDirectory() as tmp:
        src = Path(tmp) / "in.xlsx"
        src.write_bytes(workbook)
        out_dir = Path(tmp) / "out"
        subprocess.run(
            [
                soffice,
                "--headless",
                "--calc",
                "--convert-to",
                "xlsx",
                "--outdir",
                str(out_dir),
                str(src),
            ],
            check=True,
            timeout=180,
            capture_output=True,
        )
        return (out_dir / "in.xlsx").read_bytes()
