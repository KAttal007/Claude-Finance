"""Change-set builder (design doc 10.2-10.5).

Nothing is written here. The engine proposes items; the analyst accepts, edits or skips each one;
``apply.py`` writes only approved items into a new model version.
"""

from collections.abc import Callable
from dataclasses import asdict, dataclass, field
from datetime import date
from decimal import Decimal
from typing import Any, Literal, Protocol

from openpyxl.formula.translate import Translator
from openpyxl.workbook import Workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from t2r.db.models import ChangeItem, ChangeSet, Fact, Kpi
from t2r.modelsync.dependency import DependencyGraph
from t2r.modelsync.mapping import Driver, Mapping, apply_transform
from t2r.modelsync.workbook import PeriodAxis, cell_ref, find_row, fiscal_label

Kind = Literal["actual_write", "roll_forward", "assumption_suggestion", "label_shift"]


@dataclass
class SourceValue:
    value: Decimal
    fact_id: int | None
    description: str


class ValueSource(Protocol):
    def xbrl(
        self, element: str, period_start: date | None, period_end: date, consolidated: bool
    ) -> SourceValue | None: ...

    def kpi(self, name: str, period_end: date) -> SourceValue | None: ...


@dataclass
class ProposedItem:
    kind: Kind
    sheet: str
    cell: str
    old_value: str | None = None
    new_value: str | None = None
    formula_before: str | None = None
    formula_after: str | None = None
    rationale: str | None = None
    source_fact_id: int | None = None
    source: str | None = None


@dataclass
class Proposal:
    period_label: str
    items: list[ProposedItem] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)
    summary: dict[str, Any] = field(default_factory=dict)

    def as_dict(self) -> dict[str, Any]:
        return {
            "period": self.period_label,
            "summary": self.summary,
            "warnings": self.warnings,
            "items": [asdict(i) for i in self.items],
        }


# (driver, actual, current estimate) -> (suggested value, one-line rationale) or None
Suggester = Callable[[Driver, Decimal, Decimal], tuple[Decimal, str] | None]


def quarter_start(period_end: date) -> date:
    month = period_end.month - 2
    year = period_end.year
    if month <= 0:
        month += 12
        year -= 1
    return date(year, month, 1)


def build_change_set(
    wb: Workbook,
    mapping: Mapping,
    period_end: date,
    source: ValueSource,
    *,
    driver_actuals: dict[str, Decimal] | None = None,
    suggester: Suggester | None = None,
) -> Proposal:
    axis = PeriodAxis(wb, mapping)
    proposal = Proposal(period_label=fiscal_label(period_end))
    col = axis.column_for(period_end)
    if col is None:
        proposal.warnings.append(
            f"No column labelled {proposal.period_label} on {mapping.period_axis.sheet}; "
            "extend the model's period axis first"
        )
        return proposal

    p_start = quarter_start(period_end)
    written: list[tuple[str, str]] = []

    # 1. Actual writes into input cells. A formula in an estimate column is replaced by the
    #    actual only through an explicit item the analyst approves (formula_before records it).
    for row in mapping.rows:
        ws = wb[row.sheet]
        r = find_row(ws, row.row_label, mapping.period_axis.label_column)
        if r is None:
            proposal.warnings.append(f"Row {row.row_label!r} not found on {row.sheet}")
            continue
        total, fact_id, descs, missing = Decimal(0), None, [], False
        for term in row.terms():
            sv = (
                source.xbrl(term.key, p_start, period_end, mapping.consolidated)
                if term.kind == "xbrl"
                else source.kpi(term.key, period_end)
            )
            if sv is None:
                missing = True
                break
            transform = row.transform or (
                mapping.units.default_transform if term.kind == "xbrl" else "none"
            )
            total += apply_transform(sv.value, transform, term.sign)
            fact_id = fact_id or sv.fact_id
            descs.append(sv.description)
        if missing:
            proposal.warnings.append(
                f"No value for {row.row_label!r} ({row.source}) in {proposal.period_label}"
            )
            continue
        new = apply_transform(total, None, row.sign)
        cell = ws.cell(row=r, column=col)
        old = cell.value
        if isinstance(old, int | float | Decimal) and Decimal(str(old)) == new:
            continue
        is_formula = isinstance(old, str) and old.startswith("=")
        proposal.items.append(
            ProposedItem(
                kind="actual_write",
                sheet=row.sheet,
                cell=cell.coordinate,
                old_value=None if is_formula or old is None else str(old),
                new_value=str(new),
                formula_before=old if is_formula else None,
                source_fact_id=fact_id,
                source="; ".join(descs),
            )
        )
        written.append((row.sheet, cell.coordinate))

    # 2. Roll-forward: flip E -> A, extend the forecast horizon by one column.
    if axis.flag(col) == "E":
        proposal.items.append(
            ProposedItem(
                kind="roll_forward",
                sheet=axis.ws.title,
                cell=cell_ref(col, axis.flag_row),
                old_value="E",
                new_value="A",
                rationale="Period now reported",
            )
        )
    estimates = [c for c in axis.estimate_columns() if c != col]
    if estimates:
        _propose_horizon_extension(wb, axis, max(estimates), proposal)
        _list_adjacent_hardcodes(wb, mapping, axis, col, proposal)

    # 3. Assumption suggestions, one per driver, grounded in actual vs estimate.
    if suggester and driver_actuals:
        for driver in mapping.drivers:
            _propose_assumption(wb, mapping, axis, col, driver, driver_actuals, suggester, proposal)

    impact = DependencyGraph(wb).impact(written)
    counts = {
        k: sum(1 for i in proposal.items if i.kind == k)
        for k in ("actual_write", "roll_forward", "assumption_suggestion", "label_shift")
    }
    proposal.summary = {
        "writes": counts["actual_write"],
        "roll_forwards": counts["roll_forward"],
        "suggestions": counts["assumption_suggestion"],
        "label_shifts": counts["label_shift"],
        "dependents_affected": impact["dependents"],
        "unknown_impact": impact["unknown_impact"],
        "warnings": len(proposal.warnings),
    }
    return proposal


def _propose_horizon_extension(
    wb: Workbook, axis: PeriodAxis, last_col: int, proposal: Proposal
) -> None:
    """Copy the last forecast column's formulas one column right (relative refs shift)."""
    ws = axis.ws
    next_col = last_col + 1
    for r in range(1, ws.max_row + 1):
        src = ws.cell(row=r, column=last_col)
        dst = ws.cell(row=r, column=next_col)
        if dst.value not in (None, ""):
            continue
        if isinstance(src.value, str) and src.value.startswith("="):
            formula = Translator(src.value, origin=src.coordinate).translate_formula(dst.coordinate)
            proposal.items.append(
                ProposedItem(
                    kind="roll_forward",
                    sheet=ws.title,
                    cell=dst.coordinate,
                    formula_after=formula,
                    rationale=f"Extend forecast from {src.coordinate}",
                )
            )
    label = axis.label(last_col)
    proposal.items.append(
        ProposedItem(
            kind="label_shift",
            sheet=ws.title,
            cell=cell_ref(next_col, axis.header_row),
            new_value=_next_label(label),
            rationale="New forecast period",
        )
    )
    proposal.items.append(
        ProposedItem(
            kind="label_shift",
            sheet=ws.title,
            cell=cell_ref(next_col, axis.flag_row),
            new_value="E",
            rationale="New forecast period",
        )
    )


def _list_adjacent_hardcodes(
    wb: Workbook, mapping: Mapping, axis: PeriodAxis, col: int, proposal: Proposal
) -> None:
    for row in mapping.rows:
        ws = wb[row.sheet]
        r = find_row(ws, row.row_label, mapping.period_axis.label_column)
        if r is None:
            continue
        nxt = ws.cell(row=r, column=col + 1)
        if isinstance(nxt.value, int | float) and axis.flag(col + 1) == "E":
            proposal.warnings.append(
                f"Hardcoded estimate next to new actual: {row.sheet}!{nxt.coordinate} "
                f"({row.row_label}) - review"
            )


def _propose_assumption(
    wb: Workbook,
    mapping: Mapping,
    axis: PeriodAxis,
    col: int,
    driver: Driver,
    actuals: dict[str, Decimal],
    suggester: Suggester,
    proposal: Proposal,
) -> None:
    actual = actuals.get(driver.name)
    ws = wb[driver.forecast_sheet]
    r = find_row(ws, driver.forecast_row_label, mapping.period_axis.label_column)
    if actual is None or r is None:
        return
    target = ws.cell(row=r, column=col + 1)
    if not isinstance(target.value, int | float):
        return
    result = suggester(driver, actual, Decimal(str(target.value)))
    if result is None:
        return
    value, rationale = result
    proposal.items.append(
        ProposedItem(
            kind="assumption_suggestion",
            sheet=driver.forecast_sheet,
            cell=target.coordinate,
            old_value=str(target.value),
            new_value=str(value),
            rationale=rationale,
        )
    )


def _next_label(label: str) -> str:
    """Q4FY27 -> Q1FY28."""
    if len(label) >= 6 and label[0] == "Q" and label[2:4] == "FY":
        q, fy = int(label[1]), int(label[4:])
        return f"Q1FY{(fy + 1) % 100:02d}" if q == 4 else f"Q{q + 1}FY{fy:02d}"
    return f"{label}+1"


class DbValueSource:
    """Latest facts and KPIs from PostgreSQL."""

    def __init__(self, session: Session, company_id: int) -> None:
        self.session, self.company_id = session, company_id

    def xbrl(
        self, element: str, period_start: date | None, period_end: date, consolidated: bool
    ) -> SourceValue | None:
        fact = self.session.scalar(
            select(Fact).where(
                Fact.company_id == self.company_id,
                Fact.element == element,
                Fact.period_end == period_end,
                Fact.consolidated == consolidated,
                Fact.is_latest,
                Fact.dimensions == {},
                (Fact.period_start == period_start) | Fact.is_instant,
            )
        )
        if fact is None or fact.value is None:
            return None
        span = (
            f"{fact.period_start} to {fact.period_end}"
            if fact.period_start
            else f"at {fact.period_end}"
        )
        return SourceValue(fact.value, fact.id, f"{element}, {span}")

    def kpi(self, name: str, period_end: date) -> SourceValue | None:
        kpi = self.session.scalar(
            select(Kpi)
            .where(
                Kpi.company_id == self.company_id, Kpi.name == name, Kpi.period_end == period_end
            )
            .order_by(Kpi.verified.desc())
        )
        if kpi is None or kpi.value is None:
            return None
        return SourceValue(kpi.value, None, f"KPI {name}, page {kpi.page}")


def save_proposal(
    session: Session, model_id: int, base_version: int, filing_id: int | None, proposal: Proposal
) -> ChangeSet:
    cs = ChangeSet(
        model_id=model_id,
        base_version=base_version,
        filing_id=filing_id,
        summary=proposal.summary
        | {"warnings_detail": proposal.warnings, "period": proposal.period_label},
    )
    session.add(cs)
    session.flush()
    for item in proposal.items:
        session.add(
            ChangeItem(
                change_set_id=cs.id,
                kind=item.kind,
                sheet=item.sheet,
                cell=item.cell,
                old_value=item.old_value,
                new_value=item.new_value,
                formula_before=item.formula_before,
                formula_after=item.formula_after,
                rationale=item.rationale,
                source_fact_id=item.source_fact_id,
            )
        )
    session.flush()
    return cs
