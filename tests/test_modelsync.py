import io
from datetime import date
from decimal import Decimal
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from t2r.modelsync.apply import ApprovedItem, apply_headless
from t2r.modelsync.changeset import SourceValue, build_change_set, quarter_start
from t2r.modelsync.dependency import DependencyGraph, references
from t2r.modelsync.mapping import apply_transform, load_mapping, parse_source
from t2r.modelsync.workbook import fiscal_label

MAPPING = (Path(__file__).parent.parent / "mappings" / "example_quarterly.yaml").read_text()


def make_workbook() -> Workbook:
    """Data sheet: B=Q1FY27 (A), C=Q2FY27 (E, formulas), D=Q3FY27 (E). Model sheet has drivers."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Data"
    rows = [
        ["Period", "Q1FY27", "Q2FY27", "Q3FY27"],
        ["A/E", "A", "E", "E"],
        ["Revenue from operations", 90, "=B3*(1+Model!C2)", "=C3*(1+Model!D2)"],
        ["Other income", 4, 4, 4],
        ["Finance costs", 30, "=B5*1.02", "=C5*1.02"],
        ["PAT", 16.5, "=C3*0.18", "=D3*0.18"],
        ["Gold loan AUM", 80000, 82000, 84000],
        ["Margin", "=F6/C3", "=C6/C3", "=D6/D3"],
        ["Volatile", "=OFFSET(A1,1,1)", None, None],
    ]
    for r in rows:
        ws.append(r)
    model = wb.create_sheet("Model")
    model.append(["Driver", "Q1FY27", "Q2FY27", "Q3FY27"])
    model.append(["Growth", 0.1, 0.1, 0.05])
    model.append(["Yield assumption", 0.18, 0.18, 0.18])
    summary = wb.create_sheet("Summary")
    summary["A1"] = "=SUM(Data!C3:D3)"
    summary["A2"] = "=A1*2"
    return wb


class FakeSource:
    def __init__(self):
        self.facts = {
            "RevenueFromOperations": Decimal(1_000_000_000),
            "OtherIncome": Decimal(40_000_000),  # 4 crore: equal to current value, no write
            "ProfitLossForPeriod": Decimal(187_500_000),
        }

    def xbrl(self, element, period_start, period_end, consolidated):
        assert period_start == date(2026, 7, 1) and consolidated
        v = self.facts.get(element)
        return None if v is None else SourceValue(v, 11, f"{element}, Q2FY27")

    def kpi(self, name, period_end):
        return SourceValue(Decimal(83_500), None, "KPI page 7") if name == "gold_loan_aum" else None


def test_mapping_parsing_and_transforms():
    m = load_mapping(MAPPING)
    assert m.consolidated and m.period_axis.sheet == "Data"
    terms = parse_source("xbrl:A + xbrl:B - kpi:c")
    assert [(t.kind, t.key, t.sign) for t in terms] == [
        ("xbrl", "A", 1),
        ("xbrl", "B", 1),
        ("kpi", "c", -1),
    ]
    assert apply_transform(Decimal(1e7), "div:10000000") == 1
    assert apply_transform(Decimal(5), "none", -1) == -5
    with pytest.raises(ValueError):
        parse_source("bad:x")


def test_fiscal_labels():
    assert fiscal_label(date(2026, 6, 30)) == "Q1FY27"
    assert fiscal_label(date(2027, 3, 31)) == "Q4FY27"
    assert quarter_start(date(2026, 9, 30)) == date(2026, 7, 1)
    assert quarter_start(date(2027, 3, 31)) == date(2027, 1, 1)


def test_dependency_graph():
    refs, volatile = references("=SUM('Data'!C3:D3)+Model!B2", "Summary")
    assert ("Data", "C3") in refs and ("Data", "D3") in refs and ("Model", "B2") in refs
    assert not volatile
    graph = DependencyGraph(make_workbook())
    down = graph.downstream([("Data", "C3")])
    assert ("Data", "D3") in down and ("Summary", "A1") in down and ("Summary", "A2") in down
    assert "Data!B9" in graph.impact([])["unknown_impact"]


def test_change_set_proposal():
    wb = make_workbook()
    proposal = build_change_set(
        wb,
        load_mapping(MAPPING),
        date(2026, 9, 30),
        FakeSource(),
        driver_actuals={"loan_yield": Decimal("0.172")},
        suggester=lambda d, actual, est: (actual, "Yield came in lower"),
    )
    items = {(i.kind, i.sheet, i.cell): i for i in proposal.items}

    rev = items[("actual_write", "Data", "C3")]
    assert Decimal(rev.new_value) == 100 and rev.formula_before == "=B3*(1+Model!C2)"
    assert rev.source_fact_id == 11
    assert ("actual_write", "Data", "C4") not in items  # unchanged value: no item
    assert Decimal(items[("actual_write", "Data", "C7")].new_value) == 83_500  # KPI, no scaling
    assert items[("roll_forward", "Data", "C2")].new_value == "A"
    assert items[("roll_forward", "Data", "E3")].formula_after == "=D3*(1+Model!E2)"
    assert items[("label_shift", "Data", "E1")].new_value == "Q4FY27"
    assert items[("assumption_suggestion", "Model", "D3")].new_value == "0.172"

    assert any("Finance costs" in w for w in proposal.warnings)  # no fact for FinanceCosts
    assert any("Hardcoded estimate" in w and "D7" in w for w in proposal.warnings)
    assert proposal.summary["writes"] == 3
    assert proposal.summary["dependents_affected"] >= 4


def test_missing_period_column_is_a_warning_not_a_crash():
    proposal = build_change_set(
        make_workbook(), load_mapping(MAPPING), date(2027, 6, 30), FakeSource()
    )
    assert not proposal.items and "Q1FY28" in proposal.warnings[0]


def test_apply_headless_writes_only_given_items():
    wb = make_workbook()
    buf = io.BytesIO()
    wb.save(buf)
    out = apply_headless(
        buf.getvalue(),
        [
            ApprovedItem("actual_write", "Data", "C3", "100"),
            ApprovedItem("roll_forward", "Data", "C2", "A"),
            ApprovedItem("roll_forward", "Data", "E3", None, "=D3*(1+Model!E2)"),
        ],
    )
    ws = load_workbook(io.BytesIO(out))["Data"]
    assert ws["C3"].value == 100 and ws["C2"].value == "A"
    assert ws["E3"].value == "=D3*(1+Model!E2)"
    assert ws["C5"].value == "=B5*1.02"  # untouched
