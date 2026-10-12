"""End-to-end flows on PostgreSQL: start coverage -> approve assumptions -> built model ->
next results flow into Model Sync; template onboarding; model discovery. Data is SYNTHETIC."""

import io
from datetime import date
from decimal import Decimal

import pytest
from openpyxl import load_workbook
from sqlalchemy import select

from analekh.builder.coverage import BuildNotReadyError, finalize_build, start_coverage
from analekh.db.models import Assumption, Company, ModelBuild
from analekh.discovery import Matcher, WorkbookHints, discover, link
from analekh.excel_bridge.indexer import hints_for
from analekh.facts import DbValueSource
from analekh.ingest.pipeline import ingest_results_xbrl
from analekh.modelsync.changeset import build_change_set
from analekh.templates.onboarding import (
    TemplateError,
    activate,
    confirm_mapping,
    default_for,
    ensure_builtin_templates,
    upload,
)
from analekh.templates.testbuild import run_test_build
from analekh.templates_builtin import NBFC
from analekh.templates_builtin import build as build_builtin
from analekh.workbook.mapping import Mapping
from analekh.workbook.tags import Tag, read_tag, write_tag
from tests.xbrl_factory import quarter_instance

QUARTERS = [
    date(2024, 12, 31),
    date(2025, 3, 31),
    date(2025, 6, 30),
    date(2025, 9, 30),
    date(2025, 12, 31),
    date(2026, 3, 31),
    date(2026, 6, 30),
    date(2026, 9, 30),
]


def history(i: int) -> dict[str, int]:
    revenue = 800_000_000 + i * 25_000_000
    return dict(
        revenue=revenue,
        other=20_000_000,
        expenses=int(revenue * 0.7),
        tax=int((revenue + 20_000_000 - int(revenue * 0.7)) * 0.25),
    )


@pytest.fixture
def company(db, store):
    c = Company(
        nse_symbol="DEMOFIN",
        name="Demo Finance Limited",
        isin="INE000D01011",
        sector="Financial Services",
        industry="NBFC",
        results_taxonomy="indas",
    )
    db.add(c)
    db.flush()
    for i, end in enumerate(QUARTERS):
        ingest_results_xbrl(db, store, c, quarter_instance(end, **history(i)))
    return c


@pytest.fixture
def builtin(db, store):
    return {t.name: t for t in ensure_builtin_templates(db, store)}


def test_builtin_templates_registered_once(db, store, builtin):
    assert set(builtin) == {"nbfc_quarterly", "consumer_quarterly"}
    assert all(t.status == "active" and t.is_default for t in builtin.values())
    assert {t.id for t in ensure_builtin_templates(db, store)} == {t.id for t in builtin.values()}
    assert default_for(db, "NBFC").name == "nbfc_quarterly"


def test_start_coverage_to_built_model_to_quarterly_update(db, store, company, builtin):
    started = start_coverage(db, store, company, builtin["nbfc_quarterly"])
    assert started.build.status == "awaiting_assumptions"
    assert started.build.history_to == date(2026, 9, 30)
    assert len(started.fill.written) >= 8 * 4  # revenue, other income, expenses, tax, PAT ...
    assert "Finance costs" in started.fill.blank_rows  # synthetic filings do not report it
    growth = next(p for p in started.proposals if p.driver == "revenue_growth")
    assert Decimal("0.02") < growth.value < Decimal("0.04") and growth.citations
    assert "quarter-on-quarter" in growth.rationale

    with pytest.raises(BuildNotReadyError):
        finalize_build(db, store, started.model)
    for a in db.scalars(select(Assumption).where(Assumption.model_id == started.model.id)):
        a.decision, a.final_value = "accepted", a.proposed_value
    tax = db.scalar(
        select(Assumption).where(
            Assumption.model_id == started.model.id, Assumption.driver == "tax_rate"
        )
    )
    tax.decision, tax.final_value = "edited", Decimal("0.252")

    done = finalize_build(db, store, started.model)
    assert done.checks.passed, done.checks.as_dict()
    assert done.build.status == "built" and done.version is not None

    wb = load_workbook(io.BytesIO(store.get(done.version.storage_key)))
    tag = read_tag(wb)
    assert tag is not None and tag.symbol == "DEMOFIN" and tag.model_id == done.model.doc_tag
    data = wb["Data"]
    assert data["I1"].value == "Q2FY27" and data["I2"].value == "A" and data["J2"].value == "E"
    assert data["I3"].value == pytest.approx(97.5)  # revenue Q2FY27 in INR crore
    assert "RevenueFromOperations" in data["I3"].comment.text
    assert isinstance(data["J3"].value, str) and data["J3"].value.startswith("=")  # untouched
    assert wb["Drivers"]["J7"].value == pytest.approx(0.252)  # edited tax rate
    assert wb["Sources"].max_row > 30

    # Next results arrive: the built model hands off to the quarterly update with no setup.
    ingest_results_xbrl(db, store, company, quarter_instance(date(2026, 12, 31), **history(8)))
    proposal = build_change_set(
        wb,
        Mapping.model_validate(done.model.mapping),
        date(2026, 12, 31),
        DbValueSource(db, company.id),
        last_fingerprint=done.model.fingerprint,
    )
    writes = {i.cell for i in proposal.items if i.kind == "actual_write"}
    assert "J3" in writes and not any("changed outside" in w for w in proposal.warnings)
    assert any(i.kind == "roll_forward" and i.cell == "J2" for i in proposal.items)


def test_template_onboarding_from_existing_model(db, store, company):
    # An analyst's filled model for another company, uploaded as a template.
    data, _ = build_builtin(NBFC, date(2026, 9, 30))
    wb = load_workbook(io.BytesIO(data))
    for col in "BCDEFGHI":
        wb["Data"][f"{col}3"] = 50.0
    out = io.BytesIO()
    wb.save(out)

    tmpl = upload(
        db,
        store,
        out.getvalue(),
        name="my_nbfc",
        sectors=["NBFC"],
        owner="kaushal",
        from_existing_model=True,
    )
    assert tmpl.status == "draft" and tmpl.version == 1
    assert "Data!B3" in tmpl.structure["cleared"]
    suggested = Mapping.model_validate(tmpl.mapping)
    assert {r.row_label for r in suggested.rows} >= {"Revenue from operations", "Tax expense"}

    with pytest.raises(TemplateError):
        activate(db, tmpl)
    confirmed = suggested.model_dump()
    confirmed["checks"] = [
        {"sheet": "Checks", "cell_label": "All checks pass", "expect": 1, "tolerance": 0}
    ]
    confirm_mapping(db, tmpl, confirmed)
    test = run_test_build(db, store, tmpl, company)
    assert test.result == "pass", test.issues
    assert "FinanceCosts" not in test.issues["reported_items_without_row"]
    assert "Income" in test.issues["reported_items_without_row"]  # total income: formula row
    assert tmpl.status == "tested"
    assert activate(db, tmpl, make_default=True).status == "active"
    assert default_for(db, "NBFC").id == tmpl.id  # the analyst's own beats the built-in

    v2 = upload(db, store, out.getvalue(), name="my_nbfc", sectors=["NBFC"], owner="kaushal")
    assert v2.version == 2


def test_discovery_and_linking(db, store, company, tmp_path):
    matcher = Matcher([company])
    by_tag = matcher.identify(WorkbookHints("x.xlsx", tag_symbol="DEMOFIN"))
    assert (by_tag.symbol, by_tag.method) == ("DEMOFIN", "tag")
    by_name = matcher.identify(WorkbookHints("/m/DEMOFIN_Model_FY27_v3.xlsx"))
    assert (by_name.symbol, by_name.method) == ("DEMOFIN", "file_name")
    fuzzy = matcher.identify(WorkbookHints("/m/Demo Finance quarterly.xlsx"))
    assert fuzzy.symbol == "DEMOFIN" and fuzzy.confidence < 0.9
    by_isin = matcher.identify(WorkbookHints("/m/model.xlsx", text=["ISIN: INE000D01011"]))
    assert (by_isin.symbol, by_isin.method) == ("DEMOFIN", "contents")
    assert matcher.identify(WorkbookHints("/m/budget.xlsx", text=["Holiday plan"])).method == "none"

    data, _ = build_builtin(NBFC, date(2026, 9, 30))
    path = tmp_path / "budget.xlsx"
    path.write_bytes(data)
    wb = load_workbook(path)
    wb["Cover"]["B3"] = "Demo Finance Limited"
    wb.save(path)
    hints = hints_for(path)
    [candidate] = discover(db, [hints])
    assert candidate.symbol == "DEMOFIN" and candidate.method == "contents"

    model = link(db, company, hints)
    wb = load_workbook(path)
    write_tag(wb, Tag(model.doc_tag, "DEMOFIN"))
    wb.save(path)
    [again] = discover(db, [hints_for(path)])
    assert again.method == "tag" and again.model_id == model.id
    assert link(db, company, hints_for(path)).id == model.id  # relinking is idempotent
    assert db.scalar(select(ModelBuild)) is None
