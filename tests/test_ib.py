import io
from decimal import Decimal

from openpyxl import load_workbook
from pptx import Presentation
from pptx.util import Inches

from t2r.ib.comps import CompInputs, compute, summary_stats, to_xlsx
from t2r.ib.pitch import refresh
from t2r.ib.profile import ProfileData, build

CR = Decimal(10_000_000)


def comp(symbol, mcap, pat, revenue=None, debt=None, cash=None, book=None):
    return CompInputs(
        symbol,
        symbol.title(),
        mcap * CR,
        debt and debt * CR,
        cash and cash * CR,
        revenue and revenue * CR,
        None,
        pat and pat * CR,
        book and book * CR,
    )


def test_multiples_and_not_meaningful():
    row = compute(comp("AAA", 10_000, 500, revenue=4_000, debt=2_000, cash=1_000, book=5_000))
    assert row.ev == 11_000 * CR
    assert row.pe == Decimal("20.0") and row.pb == Decimal("2.0")
    assert row.ev_sales == Decimal("2.8") and row.ev_ebitda is None
    loss = compute(comp("BBB", 5_000, -100))
    assert loss.pe is None  # negative earnings: n.m., not a negative multiple


def test_summary_and_xlsx_export():
    rows = [
        compute(comp("A", 1000, 50)),
        compute(comp("B", 3000, 100)),
        compute(comp("C", 2000, -5)),
    ]
    assert summary_stats(rows)["pe"]["median"] == Decimal("25.0")
    ws = load_workbook(io.BytesIO(to_xlsx(rows)))["Comps"]
    assert ws["A5"].value == "A" and ws["H5"].value == 20.0 and ws["H7"].value == "n.m."


def test_pitch_placeholders_replaced_and_logged():
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    tf = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1)).text_frame
    para = tf.paragraphs[0]
    for part in ("PAT of {{DEMO.", "pat_q}} crore; ", "{{DEMO.unknown}}"):
        para.add_run().text = part
    buf = io.BytesIO()
    prs.save(buf)

    out, log = refresh(buf.getvalue(), lambda s, m: "187.5" if m == "pat_q" else None)
    text = Presentation(io.BytesIO(out)).slides[0].shapes[0].text_frame.text
    assert text == "PAT of 187.5 crore; {{DEMO.unknown}}"
    assert [(r.placeholder, r.value) for r in log] == [
        ("{{DEMO.pat_q}}", "187.5"),
        ("{{DEMO.unknown}}", None),
    ]


def test_profile_slide_builds():
    data = ProfileData(
        "Demo Finance",
        "DEMOFIN",
        "Gold loan NBFC.",
        [("Revenue", ["900", "1,000"]), ("PAT", ["165", "188"])],
        ["Q1FY27", "Q2FY27"],
        [("Market cap", "n.a.")],
        ["Board meeting"],
    )
    prs = Presentation(io.BytesIO(build(data)))
    assert len(prs.slides) == 1
