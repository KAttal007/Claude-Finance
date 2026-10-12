import io
from datetime import date

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from t2r.api.main import app
from t2r.db.models import Company, Model, ModelVersion
from t2r.db.session import get_session
from t2r.ingest.pipeline import ingest_results_xbrl
from t2r.modelsync.changeset import DbValueSource, build_change_set, save_proposal
from t2r.modelsync.mapping import load_mapping
from tests.test_modelsync import MAPPING, make_workbook
from tests.xbrl_factory import q1_instance, q2_instance


@pytest.fixture
def client(db, store, monkeypatch):
    monkeypatch.setattr("t2r.ingest.storage.default_store", lambda: store)
    app.dependency_overrides[get_session] = lambda: db
    yield TestClient(app)
    app.dependency_overrides.clear()


@pytest.fixture
def setup(db, store):
    company = Company(nse_symbol="DEMOFIN", name="Demo Finance Limited", is_tracked=True)
    db.add(company)
    db.flush()
    ingest_results_xbrl(db, store, company, q1_instance())
    filing = ingest_results_xbrl(db, store, company, q2_instance()).filing

    buf = io.BytesIO()
    make_workbook().save(buf)
    store.put("models/test/v1.xlsx", buf.getvalue())
    model = Model(company_id=company.id, name="demo", mapping_yaml=MAPPING)
    db.add(model)
    db.flush()
    db.add(ModelVersion(model_id=model.id, version=1, storage_key="models/test/v1.xlsx"))

    proposal = build_change_set(
        load_workbook(io.BytesIO(buf.getvalue())),
        load_mapping(MAPPING),
        date(2026, 9, 30),
        DbValueSource(db, company.id),
    )
    cs = save_proposal(db, model.id, 1, filing.id, proposal)
    return {"company": company, "filing": filing, "model": model, "cs": cs}


def test_company_endpoints(client, setup):
    assert client.get("/companies/demofin").json()["nse_symbol"] == "DEMOFIN"
    filings = client.get("/companies/DEMOFIN/filings").json()
    assert [f["period_end"] for f in filings] == ["2026-09-30", "2026-06-30"]
    facts = client.get(
        "/companies/DEMOFIN/facts",
        params={"elements": "RevenueFromOperations", "start": "2026-09-30"},
    ).json()
    assert {f["value"] for f in facts} == {"1000000000.0000", "1900000000.0000"}
    assert client.get("/companies/NOPE").status_code == 404


def test_approval_gate_then_apply_creates_new_version(client, setup, store):
    cs_id = setup["cs"].id
    cs = client.get(f"/change-sets/{cs_id}").json()
    writes = [i for i in cs["items"] if i["kind"] == "actual_write"]
    # revenue, other income (5 cr in the filing vs 4 in the model) and PAT, from DB facts
    assert {i["cell"] for i in writes} == {"C3", "C4", "C6"}

    assert client.post(f"/change-sets/{cs_id}/apply").status_code == 409  # nothing decided yet

    pat = next(i for i in writes if i["cell"] == "C6")
    r = client.post(
        f"/change-sets/{cs_id}/decisions",
        json={
            "accept_kinds": ["actual_write", "roll_forward", "label_shift"],
            "decisions": [{"item_id": pat["id"], "decision": "edited", "final_value": "18.8"}],
        },
    )
    assert r.status_code == 200
    remaining = [i for i in r.json()["items"] if i["decision"] == "pending"]
    if remaining:
        r = client.post(
            f"/change-sets/{cs_id}/decisions",
            json={"decisions": [{"item_id": i["id"], "decision": "skipped"} for i in remaining]},
        )
    assert r.json()["status"] == "approved"

    out = client.post(f"/change-sets/{cs_id}/apply").json()
    assert out["version"] == 2
    new = client.get(f"/models/{setup['model'].id}/versions/2")
    ws = load_workbook(io.BytesIO(new.content))["Data"]
    assert ws["C3"].value == 100 and ws["C6"].value == 18.8 and ws["C2"].value == "A"
    assert ws["E3"].value == "=D3*(1+Model!E2)"
    original = load_workbook(io.BytesIO(store.get("models/test/v1.xlsx")))["Data"]
    assert original["C3"].value == "=B3*(1+Model!C2)"  # original never overwritten
