import io
import uuid
from datetime import date

from openpyxl import Workbook, load_workbook

from analekh.templates_builtin import NBFC, build
from analekh.workbook.fingerprint import fingerprint
from analekh.workbook.hygiene import scan_hygiene
from analekh.workbook.labels import match_label
from analekh.workbook.periods import detect_axis, fiscal_label, parse_period
from analekh.workbook.reader import read_structure
from analekh.workbook.tags import Tag, read_tag, write_tag


def test_period_label_formats():
    q2 = date(2026, 9, 30)
    for label in ("Q2FY27", "Q2 FY27", "2QFY27", "Q2FY2027", "Sep-26", "Sep 2026", "Q2FY27E"):
        p = parse_period(label)
        assert p is not None and p.end == q2 and p.kind == "Q", label
    assert parse_period("Q2FY27E").estimate is True
    assert parse_period("Q4FY27").end == date(2027, 3, 31)
    assert parse_period("FY27E").kind == "FY" and parse_period("FY27E").end == date(2027, 3, 31)
    assert parse_period("H1FY27").end == date(2026, 9, 30)
    assert parse_period(date(2026, 12, 31)).label == "Q3FY27"
    for not_period in ("Revenue", "A/E", "2026", None, 12.5, "Apr-26"):
        assert parse_period(not_period) is None
    assert fiscal_label(date(2027, 3, 31), "FY") == "FY27"


def test_detect_axis_with_flags_and_suffixes():
    wb = Workbook()
    ws = wb.active
    ws.append(["Company XYZ"])
    ws.append(["", "Jun-26", "Sep-26", "Dec-26E", "Mar-27E"])
    ws.append(["Revenue", 1, 2, "=C3*1.1", "=D3*1.1"])
    axis = detect_axis(ws)
    assert axis is not None and axis.header_row == 2 and axis.flag_row is None
    assert axis.estimate_columns == [4, 5] and axis.last_actual_column == 3


def test_label_matching_tiers():
    assert match_label("Net sales").field == "revenue"
    assert match_label("Net Sales (INR cr)").method == "synonym"
    fuzzy = match_label("Revenue from operation")
    assert fuzzy.field == "revenue" and fuzzy.method == "fuzzy" and fuzzy.confidence < 0.98
    assert match_label("Widgets shipped to Mars").field is None
    ai = match_label("Topline", ai=lambda label, fields: ("revenue", 0.9))
    assert ai.field == "revenue" and ai.method == "ai" and ai.confidence <= 0.85


def test_structure_map_of_builtin_template():
    data, mapping = build(NBFC, date(2026, 9, 30))
    sm = read_structure(load_workbook(io.BytesIO(data)))
    assert sm.axis is not None and sm.axis.sheet == "Data"
    assert len(sm.axis.columns) == 12 and len(sm.axis.estimate_columns) == 4
    assert sm.axis.to_dict()["last_actual"] == "I"
    assert {d.label for d in sm.drivers} == {d.label for d in NBFC.drivers}
    assert {c.label for c in sm.checks} >= {"PAT difference", "All checks pass"}
    assert not sm.hygiene
    suggested = sm.suggest_mapping("nbfc_test")
    labels = {r.row_label for r in suggested.rows}
    assert "Revenue from operations" in labels
    assert "Total income" not in labels and "Profit before tax" not in labels  # formula rows
    assert labels == {r.row_label for r in mapping.rows}


def test_hygiene_findings():
    wb = Workbook()
    ws = wb.active
    ws.title = "Data"
    ws.append(["Revenue", "=B2*2", "=C2*2", 500, "=E2*2"])
    ws.append(["Base", 1, 2, 3, 4, 5])
    ws.append(["Margin", "=B2/B1", "=C2/C1", "=D2/D1+0.1", "=E2/E1"])
    ws.append(["Broken", "=#REF!*2"])
    kinds = {(f.kind, f.cell) for f in scan_hygiene(wb)}
    assert ("hardcode_in_formula_row", "D1") in kinds
    assert ("inconsistent_formula", "D3") in kinds
    assert ("error_value", "B4") in kinds


def test_fingerprint_detects_edits_and_tag_roundtrip():
    data, _ = build(NBFC, date(2026, 9, 30))
    wb = load_workbook(io.BytesIO(data))
    before = fingerprint(wb)
    wb["Data"]["B3"] = 123
    assert fingerprint(wb) != before
    tag = Tag(uuid.uuid4(), "DEMOFIN")
    write_tag(wb, tag)
    write_tag(wb, tag)
    out = io.BytesIO()
    wb.save(out)
    again = load_workbook(io.BytesIO(out.getvalue()))
    assert read_tag(again) == tag and len(again.custom_doc_props.props) == 2


def test_ai_label_resolver_is_capped_and_memoised():
    from analekh.genai.labels import label_resolver
    from analekh.workbook.labels import match_label

    calls = []

    class FakeClaude:
        small_model = "small"

        def create_json(self, purpose, schema, **kwargs):
            calls.append(kwargs["model"])
            assert None in schema["properties"]["field"]["enum"]
            return {"field": "revenue", "confidence": 0.99}

    ai = label_resolver(FakeClaude())  # type: ignore[arg-type]
    m = match_label("Topline xyz qq", ai)
    assert (m.field, m.method, m.confidence) == ("revenue", "ai", 0.85)
    match_label("Topline xyz qq", ai)
    assert calls == ["small"]
