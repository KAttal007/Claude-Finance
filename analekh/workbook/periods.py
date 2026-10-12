"""Period axis detection (design doc 12): find the header row of period labels (Q2FY27, 2QFY27,
Sep-26, FY27E, dates), the A/E row, and the last actual column.

Indian fiscal years run April to March: Q1FY27 = Apr-Jun 2026, FY27 ends 31 Mar 2027.
"""

import calendar
import re
from dataclasses import dataclass
from datetime import date, datetime

from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

_MONTHS = {m.lower(): i for i, m in enumerate(calendar.month_abbr) if m}
_Q_END_MONTH = {1: 6, 2: 9, 3: 12, 4: 3}

_PATTERNS = [
    # Q2FY27, Q2 FY27, Q2FY2027, Q2FY27E
    (re.compile(r"^Q([1-4])\s*FY\s*'?(\d{2}|\d{4})([AE])?$"), "q_fy"),
    # 2QFY27, 2Q FY27
    (re.compile(r"^([1-4])Q\s*FY\s*'?(\d{2}|\d{4})([AE])?$"), "q_fy"),
    # FY27, FY27E, FY2027A
    (re.compile(r"^FY\s*'?(\d{2}|\d{4})([AE])?$"), "fy"),
    # H1FY27, 1HFY27
    (re.compile(r"^(?:H([12])|([12])H)\s*FY\s*'?(\d{2}|\d{4})([AE])?$"), "h_fy"),
    # Sep-26, Sep 2026, Sep'26
    (re.compile(r"^([A-Za-z]{3})[\s\-']*(\d{2}|\d{4})([AE])?$"), "month"),
]


@dataclass(frozen=True)
class Period:
    kind: str  # Q, H, FY
    end: date
    estimate: bool | None  # from a trailing A/E suffix; None when the label has no suffix

    @property
    def label(self) -> str:
        return fiscal_label(self.end, self.kind)


def _year(y: str) -> int:
    n = int(y)
    return n if n > 999 else 2000 + n


def _month_end(year: int, month: int) -> date:
    return date(year, month, calendar.monthrange(year, month)[1])


def fiscal_label(period_end: date, kind: str = "Q") -> str:
    """Apr-Jun 2026 -> Q1FY27; FY ending Mar 2027 -> FY27."""
    fy = period_end.year + 1 if period_end.month >= 4 else period_end.year
    if kind == "FY":
        return f"FY{fy % 100:02d}"
    if kind == "H":
        return f"H{1 if period_end.month == 9 else 2}FY{fy % 100:02d}"
    quarter = {6: 1, 9: 2, 12: 3, 3: 4}.get(period_end.month)
    if quarter is None:
        raise ValueError(f"{period_end} is not a fiscal quarter end")
    return f"Q{quarter}FY{fy % 100:02d}"


def parse_period(value: object) -> Period | None:
    """Parse one header cell into a period, or None if it is not a period label."""
    if isinstance(value, datetime):
        value = value.date()
    if isinstance(value, date):
        if value.month in (3, 6, 9, 12) and value.day >= 28:
            return Period("Q", _month_end(value.year, value.month), None)
        return None
    if not isinstance(value, str):
        return None
    text = value.strip().upper().replace("’", "'")
    for pattern, kind in _PATTERNS:
        m = pattern.match(text)
        if not m:
            continue
        g = m.groups()
        if kind == "q_fy":
            q, fy, flag = int(g[0]), _year(g[1]), g[2]
            month = _Q_END_MONTH[q]
            year = fy if q == 4 else fy - 1
            return Period("Q", _month_end(year, month), _flag(flag))
        if kind == "fy":
            return Period("FY", date(_year(g[0]), 3, 31), _flag(g[1]))
        if kind == "h_fy":
            half = int(g[0] or g[1])
            fy = _year(g[2])
            end = _month_end(fy - 1, 9) if half == 1 else date(fy, 3, 31)
            return Period("H", end, _flag(g[3]))
        if kind == "month":
            mon = _MONTHS.get(g[0].lower()[:3])
            if mon in (3, 6, 9, 12):
                return Period("Q", _month_end(_year(g[1]), mon), _flag(g[2]))
    return None


def _flag(suffix: str | None) -> bool | None:
    return None if not suffix else suffix == "E"


@dataclass
class PeriodAxis:
    sheet: str
    header_row: int
    flag_row: int | None
    label_column: int
    columns: dict[int, Period]  # column index -> period
    flags: dict[int, str]  # column index -> "A" | "E" | ""

    def column_for(self, period_end: date, kind: str = "Q") -> int | None:
        return next(
            (c for c, p in self.columns.items() if p.end == period_end and p.kind == kind), None
        )

    def is_estimate(self, col: int) -> bool:
        flag = self.flags.get(col, "")
        if flag:
            return flag == "E"
        p = self.columns.get(col)
        return bool(p and p.estimate)

    @property
    def last_actual_column(self) -> int | None:
        actual = [c for c in self.columns if not self.is_estimate(c)]
        return max(actual) if actual else None

    @property
    def estimate_columns(self) -> list[int]:
        return sorted(c for c in self.columns if self.is_estimate(c))

    def to_dict(self) -> dict[str, object]:
        return {
            "sheet": self.sheet,
            "header_row": self.header_row,
            "flag_row": self.flag_row,
            "label_column": self.label_column,
            "columns": {
                get_column_letter(c): {
                    "label": p.label,
                    "end": p.end.isoformat(),
                    "kind": p.kind,
                    "estimate": self.is_estimate(c),
                }
                for c, p in sorted(self.columns.items())
            },
            "last_actual": get_column_letter(self.last_actual_column)
            if self.last_actual_column
            else None,
        }


def detect_axis(
    ws: Worksheet,
    max_scan_rows: int = 30,
    label_column: int | None = None,
    header_row: int | None = None,
    flag_row: int | None = None,
) -> PeriodAxis | None:
    """Find the row with the most period labels; then an adjacent row of A/E flags."""
    best: tuple[int, dict[int, Period]] | None = None
    rows = [header_row] if header_row else range(1, min(ws.max_row, max_scan_rows) + 1)
    for r in rows:
        found = {}
        for c in range(1, ws.max_column + 1):
            p = parse_period(ws.cell(row=r, column=c).value)
            if p:
                found[c] = p
        if len(found) >= 2 and (best is None or len(found) > len(best[1])):
            best = (r, found)
    if best is None:
        return None
    hdr, columns = best
    first_col = min(columns)
    label_col = label_column or max(1, first_col - 1)

    flags: dict[int, str] = {}
    candidates = [flag_row] if flag_row else [hdr + 1, hdr - 1, hdr + 2]
    found_flag_row = None
    for r in candidates:
        if r is None or r < 1:
            continue
        values = {c: str(ws.cell(row=r, column=c).value or "").strip().upper() for c in columns}
        if sum(v in ("A", "E") for v in values.values()) >= max(2, len(columns) // 2):
            flags, found_flag_row = values, r
            break
    return PeriodAxis(ws.title, hdr, found_flag_row, label_col, columns, flags)
