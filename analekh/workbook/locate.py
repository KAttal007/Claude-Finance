"""Locate mapped rows and the period axis in a workbook using a confirmed mapping."""

from datetime import date

from openpyxl import Workbook
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from analekh.workbook.mapping import Mapping
from analekh.workbook.periods import PeriodAxis, detect_axis, fiscal_label


def find_row(ws: Worksheet, label: str, label_column: int = 1) -> int | None:
    target = " ".join(label.strip().lower().split())
    for row in range(1, ws.max_row + 1):
        value = ws.cell(row=row, column=label_column).value
        if isinstance(value, str) and " ".join(value.strip().lower().split()) == target:
            return row
    return None


def cell_ref(col: int, row: int) -> str:
    return f"{get_column_letter(col)}{row}"


def axis_for(wb: Workbook, mapping: Mapping) -> PeriodAxis:
    spec = mapping.period_axis
    ws = wb[spec.sheet]
    header = find_row(ws, spec.header_row_label, spec.label_column)
    flag = find_row(ws, spec.flag_row_label, spec.label_column)
    axis = detect_axis(ws, label_column=spec.label_column, header_row=header, flag_row=flag)
    if axis is None:
        raise LookupError(f"no period axis on sheet {spec.sheet!r}")
    return axis


def relabel(wb: Workbook, axis: PeriodAxis, last_actual: date) -> PeriodAxis:
    """Rewrite the period header and A/E flags so the last actual column is ``last_actual``.

    Used when a template built for one period is filled for a company's latest quarter. Only the
    header and flag cells change; formulas are untouched.
    """
    from analekh.templates_builtin.base import quarter_ends

    cols = sorted(axis.columns)
    n_act = sum(1 for c in cols if not axis.is_estimate(c))
    ends = quarter_ends(last_actual, n_act, len(cols) - n_act)
    ws = wb[axis.sheet]
    for col, end in zip(cols, ends, strict=True):
        ws.cell(row=axis.header_row, column=col, value=fiscal_label(end))
        if axis.flag_row:
            ws.cell(row=axis.flag_row, column=col, value="A" if end <= last_actual else "E")
    new = detect_axis(
        ws, label_column=axis.label_column, header_row=axis.header_row, flag_row=axis.flag_row
    )
    assert new is not None
    return new
