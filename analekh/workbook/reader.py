"""Workbook understanding engine (design doc 12): one reader for any workbook.

It produces a structure map: period axis, input vs formula rows, labels matched to standard
fields, drivers, checks, dependencies, hygiene findings and a fingerprint. Model discovery,
template onboarding, the Coverage Builder and quarterly sync all build on it. Macros are kept
but never executed: openpyxl never runs VBA.
"""

import io
import re
from dataclasses import asdict, dataclass, field
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from analekh.workbook.dependency import DependencyGraph, references
from analekh.workbook.fingerprint import fingerprint
from analekh.workbook.hygiene import Finding, scan_hygiene
from analekh.workbook.labels import AiResolver, LabelMatch, match_label
from analekh.workbook.mapping import Check, Driver, Mapping, RowMap, Units
from analekh.workbook.mapping import PeriodAxis as AxisSpec
from analekh.workbook.periods import PeriodAxis, detect_axis

DRIVER_SHEET = re.compile(r"driver|assumption|input", re.I)
# Sheets that never hold period data to map (labels there are not matched to facts).
NON_DATA_SHEET = re.compile(r"driver|assumption|check|source|cover|index|valuation|sensitiv", re.I)
CHECK_LABEL = re.compile(r"\bcheck\b|difference|\bdiff\b|balanc", re.I)
BLUE_FONTS = {"FF0000FF", "FF0070C0", "FF0000CC", "FF1F4E79", "FF0066CC"}


@dataclass
class RowInfo:
    sheet: str
    row: int
    label: str
    kind: str  # input, formula, mixed, empty
    actual_kind: str  # kind in actual columns only
    estimate_kind: str  # kind in estimate columns only
    blue_inputs: int = 0
    match: LabelMatch | None = None


@dataclass
class DriverInfo:
    sheet: str
    row: int
    label: str
    cells: list[str]
    referenced_by: int


@dataclass
class CheckInfo:
    sheet: str
    row: int
    label: str
    cells: list[str]


@dataclass
class StructureMap:
    axis: PeriodAxis | None
    rows: list[RowInfo]
    drivers: list[DriverInfo]
    checks: list[CheckInfo]
    hygiene: list[Finding]
    fingerprint: str
    sheets: list[str]
    has_macros: bool = False
    unknown_impact: list[str] = field(default_factory=list)
    header_label: str = "Period"  # text in the label column of the period row
    flag_label: str = "A/E"

    def to_dict(self) -> dict[str, Any]:
        return {
            "sheets": self.sheets,
            "axis": self.axis.to_dict() if self.axis else None,
            "rows": [
                asdict(r) | {"match": asdict(r.match) if r.match else None} for r in self.rows
            ],
            "drivers": [asdict(d) for d in self.drivers],
            "checks": [asdict(c) for c in self.checks],
            "hygiene": [asdict(h) for h in self.hygiene],
            "fingerprint": self.fingerprint,
            "has_macros": self.has_macros,
            "unknown_impact": self.unknown_impact,
        }

    def suggest_mapping(
        self, name: str, sectors: list[str] | None = None, default_transform: str = "div:10000000"
    ) -> Mapping:
        """Suggested mapping with a confidence per row; the user confirms it once."""
        if self.axis is None:
            raise ValueError("no period axis found: cannot suggest a mapping")
        seen: set[str] = set()
        rows = []
        for r in self.rows:
            if not r.match or not r.match.source:
                continue
            # Writes go only into input cells: rows that are formulas in actual columns are
            # calculated by the workbook and never mapped.
            if r.actual_kind not in ("input", "empty"):
                continue
            if r.match.source in seen:  # first occurrence wins; duplicates go to the user
                continue
            seen.add(r.match.source)
            rows.append(
                RowMap(
                    source=r.match.source,
                    sheet=r.sheet,
                    row_label=r.label,
                    transform="none" if r.match.source.startswith("kpi:") else None,
                    confidence=r.match.confidence,
                )
            )
        drivers = []
        for d in self.drivers:
            m = match_label(d.label.replace("assumption", "").replace("growth", ""))
            drivers.append(
                Driver(
                    name=_slug(d.label),
                    actual_row_label=d.label,
                    forecast_sheet=d.sheet,
                    forecast_row_label=d.label,
                )
                if not m.field
                else Driver(
                    name=m.field,
                    actual_row_label=d.label,
                    forecast_sheet=d.sheet,
                    forecast_row_label=d.label,
                )
            )
        return Mapping(
            template=name,
            sectors=sectors or [],
            period_axis=AxisSpec(
                sheet=self.axis.sheet,
                header_row_label=self.header_label,
                flag_row_label=self.flag_label,
                label_column=self.axis.label_column,
            ),
            units=Units(default_transform=default_transform),
            rows=rows,
            drivers=drivers,
            checks=[Check(sheet=c.sheet, cell_label=c.label) for c in self.checks],
        )


def _slug(label: str) -> str:
    return re.sub(r"[^a-z0-9]+", "_", label.lower()).strip("_")


def _label_at(ws: Worksheet, row: int, col: int) -> str:
    v = ws.cell(row=row, column=col).value
    return str(v).strip() if v is not None else ""


def _cell_kind(value: object) -> str:
    if value is None or value == "":
        return "empty"
    if isinstance(value, str) and value.startswith("="):
        return "formula"
    if isinstance(value, int | float):
        return "input"
    return "text"


def _combine(kinds: list[str]) -> str:
    kinds = [k for k in kinds if k in ("input", "formula")]
    if not kinds:
        return "empty"
    return kinds[0] if len(set(kinds)) == 1 else "mixed"


def has_macros(data: bytes) -> bool:
    return b"vbaProject.bin" in data


def load(data: bytes) -> Workbook:
    """Load bytes; macros are carried along (never executed) only when the file has them."""
    return load_workbook(io.BytesIO(data), keep_vba=has_macros(data))


def read_structure(
    wb: Workbook, ai: AiResolver | None = None, axis_sheet: str | None = None
) -> StructureMap:
    sheets = wb.sheetnames
    axis: PeriodAxis | None = None
    candidates = [wb[axis_sheet]] if axis_sheet else wb.worksheets
    for ws in candidates:
        found = detect_axis(ws)
        if found and (axis is None or len(found.columns) > len(axis.columns)):
            axis = found

    graph = DependencyGraph(wb)
    rows: list[RowInfo] = []
    for ws in wb.worksheets:
        # Period-aligned sheets share the axis columns (same column = same period).
        aligned = axis is not None and (
            ws.title == axis.sheet or not NON_DATA_SHEET.search(ws.title)
        )
        label_col = axis.label_column if axis and aligned else 1
        data_cols = (
            list(axis.columns)
            if axis and aligned
            else list(range(label_col + 1, ws.max_column + 1))
        )
        est = set(axis.estimate_columns) if axis and aligned else set()
        for r in range(1, ws.max_row + 1):
            if axis and ws.title == axis.sheet and r in (axis.header_row, axis.flag_row):
                continue
            label = _label_at(ws, r, label_col)
            if not label:
                continue
            cells = [ws.cell(row=r, column=c) for c in data_cols]
            kinds = [_cell_kind(c.value) for c in cells]
            info = RowInfo(
                sheet=ws.title,
                row=r,
                label=label,
                kind=_combine(kinds),
                actual_kind=_combine(
                    [k for c, k in zip(data_cols, kinds, strict=True) if c not in est]
                ),
                estimate_kind=_combine(
                    [k for c, k in zip(data_cols, kinds, strict=True) if c in est]
                ),
                blue_inputs=sum(
                    1
                    for c in cells
                    if _cell_kind(c.value) == "input"
                    and c.font is not None
                    and c.font.color is not None
                    and str(c.font.color.rgb) in BLUE_FONTS
                ),
                match=match_label(label, ai) if aligned else None,
            )
            rows.append(info)

    drivers = _find_drivers(wb, graph)
    checks = _find_checks(wb)
    header_label, flag_label = "Period", "A/E"
    if axis:
        ws = wb[axis.sheet]
        header_label = _label_at(ws, axis.header_row, axis.label_column) or header_label
        if axis.flag_row:
            flag_label = _label_at(ws, axis.flag_row, axis.label_column) or flag_label
    return StructureMap(
        axis=axis,
        rows=rows,
        drivers=drivers,
        checks=checks,
        hygiene=scan_hygiene(wb, axis),
        fingerprint=fingerprint(wb),
        sheets=sheets,
        has_macros=getattr(wb, "vba_archive", None) is not None,
        unknown_impact=sorted(f"{s}!{c}" for s, c in graph.volatile),
        header_label=header_label,
        flag_label=flag_label,
    )


def _find_drivers(wb: Workbook, graph: DependencyGraph) -> list[DriverInfo]:
    """Input cells on assumption sheets that forecast formulas on other sheets reference."""
    drivers = []
    for ws in wb.worksheets:
        if not DRIVER_SHEET.search(ws.title):
            continue
        for r in range(1, ws.max_row + 1):
            label = _label_at(ws, r, 1)
            if not label:
                continue
            cells = [c for c in ws[r][1:] if _cell_kind(c.value) == "input"]
            if not cells:
                continue
            referencing = sum(
                1
                for c in cells
                for s, _ in graph.dependents.get((ws.title, c.coordinate), ())
                if s != ws.title
            )
            if referencing:
                drivers.append(
                    DriverInfo(ws.title, r, label, [c.coordinate for c in cells], referencing)
                )
    return drivers


def _find_checks(wb: Workbook) -> list[CheckInfo]:
    """Rows of formulas on a Checks sheet, or rows labelled as checks/differences elsewhere."""
    checks = []
    for ws in wb.worksheets:
        on_checks_sheet = "check" in ws.title.lower()
        for r in range(1, ws.max_row + 1):
            label = _label_at(ws, r, 1)
            if not label or label.lower().startswith("period"):
                continue
            if not (on_checks_sheet or CHECK_LABEL.search(label)):
                continue
            cells = [
                c.coordinate
                for c in ws[r][1:]
                if isinstance(c.value, str)
                and c.value.startswith("=")
                and references(c.value, ws.title)[0]
            ]
            if cells:
                checks.append(CheckInfo(ws.title, r, label, cells))
    return checks
