"""Pitch-page refresh (design doc 12): replace ``{{SYMBOL.metric}}`` placeholders in a PowerPoint
file with current values and log every replacement.

Placeholders may be split across text runs by PowerPoint; replacement works on whole paragraphs
and rewrites the first run, which keeps that run's formatting.
"""

import io
import re
from collections.abc import Callable
from dataclasses import dataclass

from pptx import Presentation

PLACEHOLDER = re.compile(r"\{\{\s*([A-Z0-9&_-]+)\.([a-z0-9_]+)\s*\}\}")


@dataclass
class Replacement:
    slide: int
    placeholder: str
    value: str | None  # None = no value available, left in place


def refresh(
    pptx: bytes, lookup: Callable[[str, str], str | None]
) -> tuple[bytes, list[Replacement]]:
    prs = Presentation(io.BytesIO(pptx))
    log: list[Replacement] = []
    for n, slide in enumerate(prs.slides, start=1):
        for shape in slide.shapes:
            frames = []
            if shape.has_text_frame:
                frames.append(shape.text_frame)
            if getattr(shape, "has_table", False) and shape.has_table:
                frames += [cell.text_frame for row in shape.table.rows for cell in row.cells]
            for tf in frames:
                for para in tf.paragraphs:
                    text = "".join(r.text for r in para.runs)
                    if not PLACEHOLDER.search(text):
                        continue

                    def sub(m: re.Match[str], slide_no: int = n) -> str:
                        value = lookup(m.group(1), m.group(2))
                        log.append(Replacement(slide_no, m.group(0), value))
                        return value if value is not None else m.group(0)

                    new = PLACEHOLDER.sub(sub, text)
                    para.runs[0].text = new
                    for r in para.runs[1:]:
                        r.text = ""
    out = io.BytesIO()
    prs.save(out)
    return out.getvalue(), log
