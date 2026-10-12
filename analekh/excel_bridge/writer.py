"""Write approved changes into the analyst's workbook (design doc 13.3, 16.4-16.5).

xlwings drives desktop Excel when it is available (native recalculation, formatting kept, the
open workbook is used directly); otherwise openpyxl writes the file. Save modes: a new
versioned copy next to the original (default), or in place after an automatic backup. The
original is never silently overwritten.
"""

import shutil
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook

from analekh.modelsync.apply import coerce
from analekh.workbook.tags import Tag, write_tag


@dataclass
class Write:
    sheet: str
    cell: str
    value: str | None
    formula: str | None = None


def target_path(src: Path, save_mode: str, label: str) -> Path:
    stamp = datetime.now().strftime("%Y%m%d-%H%M%S")
    if save_mode == "in_place":
        backup = src.with_name(f"{src.stem}.backup-{stamp}{src.suffix}")
        shutil.copy2(src, backup)
        return src
    return src.with_name(f"{src.stem}.{label}-{stamp}{src.suffix}")


def apply_openpyxl(
    src: Path, writes: list[Write], save_mode: str, label: str, tag: Tag | None = None
) -> Path:
    wb = load_workbook(src, keep_vba=src.suffix.lower() == ".xlsm")
    for w in writes:
        wb[w.sheet][w.cell].value = w.formula if w.formula else coerce(w.value)
    if tag:
        write_tag(wb, tag)
    wb.calculation.fullCalcOnLoad = True
    dst = target_path(src, save_mode, label)
    wb.save(dst)
    return dst


def apply_xlwings(
    src: Path, writes: list[Write], save_mode: str, label: str
) -> Path:  # pragma: no cover - needs desktop Excel
    import xlwings as xw

    dst = target_path(src, save_mode, label)
    book = next(
        (b for app in xw.apps for b in app.books if Path(b.fullname).resolve() == src.resolve()),
        None,
    )
    opened_here = book is None
    app = xw.App(visible=False, add_book=False) if opened_here else None
    try:
        book = book or app.books.open(str(src))  # type: ignore[union-attr]
        for w in writes:
            rng = book.sheets[w.sheet].range(w.cell)
            if w.formula:
                rng.formula = w.formula
            else:
                rng.value = coerce(w.value)
        book.app.calculate()
        book.save(str(dst))
        if opened_here:
            book.close()
    finally:
        if app is not None:
            app.quit()
    return dst


def apply(src: Path, writes: list[Write], save_mode: str, label: str) -> Path:
    try:
        import xlwings  # noqa: F401
    except ImportError:
        return apply_openpyxl(src, writes, save_mode, label)
    try:
        return apply_xlwings(src, writes, save_mode, label)
    except Exception:  # noqa: BLE001 - no Excel on this machine: headless fallback
        return apply_openpyxl(src, writes, save_mode, label)


def link_tag(path: Path, tag: Tag) -> None:
    """Write the hidden tag after the user confirms a link (custom document properties)."""
    wb = load_workbook(path, keep_vba=path.suffix.lower() == ".xlsm")
    write_tag(wb, tag)
    wb.save(path)
