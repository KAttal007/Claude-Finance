"""Find models in watched folders and extract identification hints (design doc 13.1-13.2).

Runs on the analyst's machine. Reads .xlsx/.xlsm with openpyxl (macros are never executed),
collects the hidden tag, header and cover text, modified time and fingerprint, and sends only
those hints to the server: workbooks themselves stay on the user's machine (section 23).
"""

from collections.abc import Iterator
from pathlib import Path

from openpyxl import load_workbook

from analekh.discovery import WorkbookHints
from analekh.workbook.fingerprint import fingerprint
from analekh.workbook.tags import read_tag

EXTENSIONS = {".xlsx", ".xlsm"}
SCAN_ROWS, SCAN_COLS, SCAN_SHEETS = 30, 10, 3


def iter_workbooks(folders: list[Path]) -> Iterator[Path]:
    for folder in folders:
        if not folder.exists():
            continue
        for path in folder.rglob("*"):
            if path.suffix.lower() in EXTENSIONS and not path.name.startswith("~$"):
                yield path


def hints_for(path: Path) -> WorkbookHints:
    wb = load_workbook(
        path, read_only=False, keep_vba=path.suffix.lower() == ".xlsm", data_only=True
    )
    tag = read_tag(wb)
    text: list[str] = []
    for ws in wb.worksheets[:SCAN_SHEETS]:
        for row in ws.iter_rows(min_row=1, max_row=SCAN_ROWS, max_col=SCAN_COLS):
            for c in row:
                if isinstance(c.value, str) and 2 <= len(c.value.strip()) <= 120:
                    text.append(c.value.strip())
    return WorkbookHints(
        path=str(path),
        tag_model_id=str(tag.model_id) if tag else None,
        tag_symbol=tag.symbol if tag else None,
        text=text[:400],
        modified=path.stat().st_mtime,
        fingerprint=fingerprint(wb),
    )


def scan(folders: list[Path]) -> list[WorkbookHints]:
    out = []
    for path in iter_workbooks(folders):
        try:
            out.append(hints_for(path))
        except Exception:  # noqa: BLE001 - unreadable or locked file: skip, keep scanning
            continue
    return out
