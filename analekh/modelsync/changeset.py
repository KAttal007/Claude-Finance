"""Change-set builder (design doc 16).

Uses the workbook understanding engine to locate periods and rows. If the workbook's fingerprint
differs from the last one seen, the structure is rescanned and the change set says so (6.4).
Nothing is written here. The engine proposes items; the analyst accepts, edits or skips each one;
``apply.py`` writes only approved items into a new model version.
"""

from collections.abc import Callable
from dataclasses import asdict, dataclass, field
from datetime import date
from decimal import Decimal
from typing import Any, Literal

from openpyxl.formula.translate import Translator
from openpyxl.workbook import Workbook
from sqlalchemy.orm import Session

from analekh.db.models import ChangeItem, ChangeSet
from analekh.facts import DbValueSource, SourceValue, ValueSource, quarter_start, resolve_row
from analekh.workbook.dependency import DependencyGraph
from analekh.workbook.fingerprint import fingerprint
from analekh.workbook.locate import axis_for, cell_ref, find_row
from analekh.workbook.mapping import Driver, Mapping
from analekh.workbook.periods import PeriodAxis, fiscal_label

__all__ = [
    "DbValueSource",
    "Proposal",
    "ProposedItem",
    "SourceValue",
    "ValueSource",
    "build_change_set",
    "quarter_start",
    "save_proposal",
]

Kind = Literal["actual_write", "roll_forward", "assumption_suggestion", "label_shift"]


@dataclass
class ProposedItem:
    kind: Kind
    sheet: str
    cell: str
    old_value: str | None = None
    new_value: str | None = None
    formula_before: str | None = None
    formula_after: str | None = None
    rationale: str | None = None
    source_fact_id: int | None = None
    source: str | None = None


@dataclass
class Proposal:
    period_label: str
    items: list[ProposedItem] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)
    summary: dict[str, Any] = field(default_factory=dict)

    def as_dict(self) -> dict[str, Any]:
        return {
            "period": self.period_label,
            "summary": self.summary,
            "warnings": self.warnings,
            "items": [asdict(i) for i in self.items],
        }


# (driver, actual, current estimate) -> (suggested value, one-line rationale) or None
Suggester = Callable[[Driver, Decimal, Decimal], tuple[Decimal, str] | None]


def build_change_set(
    wb: Workbook,
    mapping: Mapping,
    period_end: date,
    source: ValueSource,
    *,
    driver_actuals: dict[str, Decimal] | None = None,
    suggester: Suggester | None = None,
    last_fingerprint: str | None = None,
) -> Proposal:
    proposal = Proposal(period_label=fiscal_label(period_end))
    if last_fingerprint and fingerprint(wb) != last_fingerprint:
        proposal.warnings.append(
            "Workbook changed outside Analekh since the last sync; structure and dependencies "
            "were rescanned before planning"
        )
    axis = axis_for(wb, mapping)
    col = axis.column_for(period_end)
    if col is None:
        proposal.warnings.append(
            f"No column labelled {proposal.period_label} on {mapping.period_axis.sheet}; "
            "extend the model's period axis first"
        )
        return proposal

    written: list[tuple[str, str]] = []

    # 1. Actual writes into input cells. A formula in an estimate column is replaced by the
    #    actual only through an explicit item the analyst approves (formula_before records it).
    for row in mapping.rows:
        ws = wb[row.sheet]
        r = find_row(ws, row.row_label, mapping.period_axis.label_column)
        if r is None:
            proposal.warnings.append(f"Row {row.row_label!r} not found on {row.sheet}")
            continue
        resolved = resolve_row(source, row, mapping, period_end)
        if resolved is None:
            proposal.warnings.append(
                f"No value for {row.row_label!r} ({row.source}) in {proposal.period_label}"
            )
            continue
        new = resolved.value
        if row.confidence < 0.8:
            proposal.warnings.append(
                f"Low-confidence mapping for {row.row_label!r} ({row.confidence:.2f}): check the "
                "source side by side"
            )
        cell = ws.cell(row=r, column=col)
        old = cell.value
        if isinstance(old, int | float | Decimal) and Decimal(str(old)) == new:
            continue
        is_formula = isinstance(old, str) and old.startswith("=")
        proposal.items.append(
            ProposedItem(
                kind="actual_write",
                sheet=row.sheet,
                cell=cell.coordinate,
                old_value=None if is_formula or old is None else str(old),
                new_value=str(new),
                formula_before=old if is_formula else None,
                source_fact_id=resolved.fact_id,
                source="; ".join(resolved.descriptions),
            )
        )
        written.append((row.sheet, cell.coordinate))

    # 2. Roll-forward: flip E -> A, extend the forecast horizon by one column.
    if axis.is_estimate(col) and axis.flag_row:
        proposal.items.append(
            ProposedItem(
                kind="roll_forward",
                sheet=axis.sheet,
                cell=cell_ref(col, axis.flag_row),
                old_value="E",
                new_value="A",
                rationale="Period now reported",
            )
        )
    estimates = [c for c in axis.estimate_columns if c != col]
    if estimates:
        _propose_horizon_extension(wb, axis, max(estimates), proposal)
        _list_adjacent_hardcodes(wb, mapping, axis, col, proposal)

    # 3. Assumption suggestions, one per driver, grounded in actual vs estimate.
    if suggester and driver_actuals:
        for driver in mapping.drivers:
            _propose_assumption(wb, mapping, axis, col, driver, driver_actuals, suggester, proposal)

    impact = DependencyGraph(wb).impact(written)
    counts = {
        k: sum(1 for i in proposal.items if i.kind == k)
        for k in ("actual_write", "roll_forward", "assumption_suggestion", "label_shift")
    }
    proposal.summary = {
        "writes": counts["actual_write"],
        "roll_forwards": counts["roll_forward"],
        "suggestions": counts["assumption_suggestion"],
        "label_shifts": counts["label_shift"],
        "dependents_affected": impact["dependents"],
        "unknown_impact": impact["unknown_impact"],
        "warnings": len(proposal.warnings),
    }
    return proposal


def _propose_horizon_extension(
    wb: Workbook, axis: PeriodAxis, last_col: int, proposal: Proposal
) -> None:
    """Copy the last forecast column's formulas one column right (relative refs shift)."""
    ws = wb[axis.sheet]
    next_col = last_col + 1
    for r in range(1, ws.max_row + 1):
        src = ws.cell(row=r, column=last_col)
        dst = ws.cell(row=r, column=next_col)
        if dst.value not in (None, ""):
            continue
        if isinstance(src.value, str) and src.value.startswith("="):
            formula = Translator(src.value, origin=src.coordinate).translate_formula(dst.coordinate)
            proposal.items.append(
                ProposedItem(
                    kind="roll_forward",
                    sheet=ws.title,
                    cell=dst.coordinate,
                    formula_after=formula,
                    rationale=f"Extend forecast from {src.coordinate}",
                )
            )
    label = axis.columns[last_col].label
    proposal.items.append(
        ProposedItem(
            kind="label_shift",
            sheet=ws.title,
            cell=cell_ref(next_col, axis.header_row),
            new_value=_next_label(label),
            rationale="New forecast period",
        )
    )
    if axis.flag_row:
        proposal.items.append(
            ProposedItem(
                kind="label_shift",
                sheet=ws.title,
                cell=cell_ref(next_col, axis.flag_row),
                new_value="E",
                rationale="New forecast period",
            )
        )


def _list_adjacent_hardcodes(
    wb: Workbook, mapping: Mapping, axis: PeriodAxis, col: int, proposal: Proposal
) -> None:
    for row in mapping.rows:
        ws = wb[row.sheet]
        r = find_row(ws, row.row_label, mapping.period_axis.label_column)
        if r is None:
            continue
        nxt = ws.cell(row=r, column=col + 1)
        if isinstance(nxt.value, int | float) and axis.is_estimate(col + 1):
            proposal.warnings.append(
                f"Hardcoded estimate next to new actual: {row.sheet}!{nxt.coordinate} "
                f"({row.row_label}) - review"
            )


def _propose_assumption(
    wb: Workbook,
    mapping: Mapping,
    axis: PeriodAxis,
    col: int,
    driver: Driver,
    actuals: dict[str, Decimal],
    suggester: Suggester,
    proposal: Proposal,
) -> None:
    actual = actuals.get(driver.name)
    ws = wb[driver.forecast_sheet]
    r = find_row(ws, driver.forecast_row_label, 1)
    if actual is None or r is None:
        return
    target = ws.cell(row=r, column=col + 1)
    if not isinstance(target.value, int | float):
        return
    result = suggester(driver, actual, Decimal(str(target.value)))
    if result is None:
        return
    value, rationale = result
    proposal.items.append(
        ProposedItem(
            kind="assumption_suggestion",
            sheet=driver.forecast_sheet,
            cell=target.coordinate,
            old_value=str(target.value),
            new_value=str(value),
            rationale=rationale,
        )
    )


def _next_label(label: str) -> str:
    """Q4FY27 -> Q1FY28."""
    if len(label) >= 6 and label[0] == "Q" and label[2:4] == "FY":
        q, fy = int(label[1]), int(label[4:])
        return f"Q1FY{(fy + 1) % 100:02d}" if q == 4 else f"Q{q + 1}FY{fy:02d}"
    return f"{label}+1"


def save_proposal(
    session: Session, model_id: int, base_version: int, filing_id: int | None, proposal: Proposal
) -> ChangeSet:
    cs = ChangeSet(
        model_id=model_id,
        base_version=base_version,
        filing_id=filing_id,
        summary=proposal.summary
        | {"warnings_detail": proposal.warnings, "period": proposal.period_label},
    )
    session.add(cs)
    session.flush()
    for item in proposal.items:
        session.add(
            ChangeItem(
                change_set_id=cs.id,
                kind=item.kind,
                sheet=item.sheet,
                cell=item.cell,
                old_value=item.old_value,
                new_value=item.new_value,
                formula_before=item.formula_before,
                formula_after=item.formula_after,
                rationale=item.rationale,
                source_fact_id=item.source_fact_id,
            )
        )
    session.flush()
    return cs
