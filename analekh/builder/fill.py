"""Fill a template's actual columns with sourced history (design doc 15.1 step 3, 14.3).

Inputs are filled, formulas are never touched, every written cell gets a comment naming its
source, and the Sources sheet lists every input with document and page. Shared by the
Coverage Builder and template test builds.
"""

from dataclasses import dataclass, field
from datetime import date

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Font

from analekh.facts import ValueSource, resolve_row
from analekh.workbook.locate import axis_for, find_row, relabel
from analekh.workbook.mapping import Mapping
from analekh.workbook.periods import PeriodAxis

BLUE = Font(color="FF0000FF")
AUTHOR = "Analekh"


@dataclass
class Written:
    sheet: str
    cell: str
    value: str
    source: str
    period: str
    filing_id: int | None
    page: int | None


@dataclass
class FillReport:
    axis: PeriodAxis
    written: list[Written] = field(default_factory=list)
    blank_rows: list[str] = field(default_factory=list)  # template rows the company never reports
    partial_rows: dict[str, list[str]] = field(default_factory=dict)  # row -> missing periods
    missing_rows: list[str] = field(default_factory=list)  # mapped label not found in workbook
    skipped_formula_cells: list[str] = field(default_factory=list)

    def as_dict(self) -> dict[str, object]:
        return {
            "written": len(self.written),
            "blank_rows": self.blank_rows,
            "partial_rows": self.partial_rows,
            "missing_rows": self.missing_rows,
            "skipped_formula_cells": self.skipped_formula_cells,
        }


def fill_history(
    wb: Workbook, mapping: Mapping, source: ValueSource, last_actual: date
) -> FillReport:
    axis = relabel(wb, axis_for(wb, mapping), last_actual)
    report = FillReport(axis=axis)
    actual_cols = [c for c in sorted(axis.columns) if not axis.is_estimate(c)]

    for row in mapping.rows:
        ws = wb[row.sheet]
        r = find_row(ws, row.row_label, mapping.period_axis.label_column)
        if r is None:
            report.missing_rows.append(f"{row.sheet}!{row.row_label}")
            continue
        missing_periods = []
        for col in actual_cols:
            period = axis.columns[col]
            cell = ws.cell(row=r, column=col)
            if isinstance(cell.value, str) and cell.value.startswith("="):
                report.skipped_formula_cells.append(f"{row.sheet}!{cell.coordinate}")
                continue
            resolved = resolve_row(source, row, mapping, period.end)
            if resolved is None:
                missing_periods.append(period.label)
                continue
            cell.value = float(resolved.value)
            cell.font = BLUE
            desc = "; ".join(resolved.descriptions)
            cell.comment = Comment(f"Source: {desc}", AUTHOR)
            first = resolved.parts[0]
            report.written.append(
                Written(
                    row.sheet,
                    cell.coordinate,
                    str(resolved.value),
                    desc,
                    period.label,
                    first.filing_id,
                    first.page,
                )
            )
        if len(missing_periods) == len(actual_cols):
            report.blank_rows.append(row.row_label)
        elif missing_periods:
            report.partial_rows[row.row_label] = missing_periods
    write_sources(wb, report.written)
    return report


def write_sources(wb: Workbook, written: list[Written]) -> None:
    if "Sources" not in wb.sheetnames:
        ws = wb.create_sheet("Sources")
        ws.append(["Sheet", "Cell", "Value", "Source", "Period", "Filing", "Document and page"])
    ws = wb["Sources"]
    for w in written:
        ws.append(
            [
                w.sheet,
                w.cell,
                w.value,
                w.source,
                w.period,
                w.filing_id,
                f"page {w.page}" if w.page else "XBRL",
            ]
        )


def write_drivers(
    wb: Workbook, mapping: Mapping, axis: PeriodAxis, values: dict[str, float]
) -> list[str]:
    """Write accepted driver values into every estimate column of each driver row."""
    written = []
    by_name = {d.name: d for d in mapping.drivers}
    for name, value in values.items():
        d = by_name.get(name)
        if d is None:
            continue
        ws = wb[d.forecast_sheet]
        r = find_row(ws, d.forecast_row_label, 1)
        if r is None:
            continue
        for col in axis.estimate_columns:
            cell = ws.cell(row=r, column=col)
            if isinstance(cell.value, str) and cell.value.startswith("="):
                continue  # never overwrite a formula
            cell.value = value
            cell.font = BLUE
            written.append(f"{ws.title}!{cell.coordinate}")
    return written
